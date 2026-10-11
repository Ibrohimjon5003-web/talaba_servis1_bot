import asyncio
import json
import logging
import os
import re
from pathlib import Path

from aiogram import Bot, Dispatcher, F, types
from aiogram.filters import CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import (
    BufferedInputFile,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    KeyboardButton,
    ReplyKeyboardMarkup,
)
from dotenv import load_dotenv
from docx import Document as WordDocument
from openai import OpenAI
from pptx import Presentation

load_dotenv()

BOT_TOKEN = os.getenv("BOT_TOKEN", "")
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY", "")
OUTPUT_DIR = Path("output")
OUTPUT_DIR.mkdir(exist_ok=True)

if not BOT_TOKEN:
    raise ValueError("BOT_TOKEN topilmadi. .env faylni to'ldiring.")

bot = Bot(token=BOT_TOKEN)
dp = Dispatcher()

main_menu = ReplyKeyboardMarkup(
    keyboard=[
        [KeyboardButton(text="📚 Xizmatlar"), KeyboardButton(text="ℹ️ Biz haqimizda")],
        [KeyboardButton(text="✍️ Buyurtma berish"), KeyboardButton(text="📞 Bog'lanish")],
    ],
    resize_keyboard=True,
)

service_menu = InlineKeyboardMarkup(
    inline_keyboard=[
        [
            InlineKeyboardButton(text="📊 Slayd", callback_data="service_slayd"),
            InlineKeyboardButton(text="📄 Referat", callback_data="service_referat"),
        ],
        [
            InlineKeyboardButton(text="📰 Maqola / Tezis", callback_data="service_maqola"),
            InlineKeyboardButton(text="🧩 Krossvord / Test", callback_data="service_test"),
        ],
        [
            InlineKeyboardButton(text="✍️ Essay / Mustaqil ish", callback_data="service_essay"),
        ],
    ]
)


class OrderForm(StatesGroup):
    waiting_for_service = State()
    waiting_for_topic = State()
    waiting_for_slide_count = State()
    waiting_for_language = State()
    generating = State()


SERVICES = {
    "slayd": "📊 Slayd/Taqdimot",
    "referat": "📄 Referat",
    "maqola": "📰 Maqola / Tezis",
    "test": "🧩 Krossvord / Test",
    "essay": "✍️ Essay / Mustaqil ish",
}


def normalize_service_name(raw: str) -> str:
    return raw.strip().lower().replace("service_", "")


def parse_slide_count(text: str):
    text = text.strip()
    if not text.isdigit():
        return None
    number = int(text)
    if number < 3 or number > 30:
        return None
    return number


def clean_json_text(text: str) -> str:
    text = text.strip()
    match = re.search(r"```(?:json)?\s*(.*?)\s*```", text, re.S)
    if match:
        return match.group(1).strip()
    return text


def build_prompt(topic: str, service: str, slide_count: int, language: str) -> str:
    service_label = SERVICES.get(service, "Umumiy ish")
    return f"""
    Siz - yuqori malakali akademik yozuvchi va taqdimot muharriri.

    Quyidagi talab asosida sifatli va zamonaviy {service_label} tayyorlang.

    Ma'lumotlar:
    - Mavzu: {topic}
    - Slaydlar soni: {slide_count}
    - Til: {language}

    Talablar:
    1. Faqatgina JSON formatda javob bering.
    2. JSON tarkibida quyidagi strukturani saqlang:
    {
      "title": "Keng sarlavha",
      "subtitle": "Kichik izoh",
      "slides": [
        {"title": "1-slayd sarlavhasi", "bullets": ["b1", "b2", "b3"]},
        {"title": "2-slayd sarlavhasi", "bullets": ["b1", "b2"]}
      ]
    }
    3. Sarlavha va slaydlarni akademik, tushunarli va ilmiy uslubda yozing.
    4. Slaydlar soni {slide_count} bo'lsin, har bir slaydda 3-5 ta nuqta bo'lsin.
    5. Tilga mos ravishda matnlarni yozing.
    6. Faqat JSON matni, boshqa izohlarni yozmang.
    """


async def generate_openai_content(topic: str, service: str, slide_count: int, language: str) -> dict:
    if not OPENAI_API_KEY:
        raise ValueError("OPENAI_API_KEY topilmadi. .env faylni to'ldiring.")

    client = OpenAI(api_key=OPENAI_API_KEY)
    prompt = build_prompt(topic, service, slide_count, language)

    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "user", "content": prompt}],
        temperature=0.7,
        max_tokens=2000,
    )

    raw = response.choices[0].message.content
    if not raw:
        raise ValueError("OpenAI javobi bo'sh qaytdi.")

    cleaned = clean_json_text(raw)
    try:
        data = json.loads(cleaned)
    except json.JSONDecodeError as exc:
        raise ValueError(f"OpenAI JSON not valid: {exc}") from exc

    if not isinstance(data, dict) or "slides" not in data:
        raise ValueError("OpenAI JSON format noto'g'ri qaytadi.")

    return data


def create_pptx_document(content: dict, filename: str) -> Path:
    prs = Presentation()
    prs.slide_width = 12192000
    prs.slide_height = 6858000

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = content.get("title", "Academic Presentation")
    title_slide.placeholders[1].text = content.get("subtitle", "Generated by AI")

    for slide_data in content.get("slides", []):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = slide_data.get("title", "Slide")
        body = slide.shapes.placeholders[1].text_frame
        body.clear()
        for idx, bullet in enumerate(slide_data.get("bullets", [])):
            p = body.paragraphs[0] if idx == 0 else body.add_paragraph()
            p.text = bullet
            p.level = 0

    path = OUTPUT_DIR / filename
    prs.save(path)
    return path


def create_word_document(content: dict, filename: str) -> Path:
    doc = WordDocument()
    doc.add_heading(content.get("title", "Academic Work"), level=0)
    doc.add_paragraph(content.get("subtitle", "Generated by AI"))

    for index, slide_data in enumerate(content.get("slides", []), start=1):
        doc.add_heading(f"{index}. {slide_data.get('title', 'Slide')}", level=1)
        for bullet in slide_data.get("bullets", []):
            doc.add_paragraph(f"• {bullet}")

    path = OUTPUT_DIR / filename
    doc.save(path)
    return path


async def send_generated_files(chat_id: int, topic: str, service: str, slide_count: int, language: str):
    try:
        generated = await asyncio.to_thread(generate_openai_content, topic, service, slide_count, language)
    except Exception as exc:
        await bot.send_message(chat_id, f"⚠️ AI generatsiyasi vaqtida xatolik yuz berdi:\n{exc}")
        return

    safe_topic = re.sub(r"[^a-zA-Z0-9_\-]+", "_", topic.strip())[:30] or "tema"
    file_name = f"{service}_{safe_topic}_{slide_count}_{language}.pptx"
    doc_name = f"{service}_{safe_topic}_{slide_count}_{language}.docx"

    ppt_path = create_pptx_document(generated, file_name)
    doc_path = create_word_document(generated, doc_name)

    await bot.send_document(
        chat_id=chat_id,
        document=BufferedInputFile(ppt_path.read_bytes(), ppt_path.name),
        caption="📊 Tayyor PowerPoint fayli",
    )

    await bot.send_document(
        chat_id=chat_id,
        document=BufferedInputFile(doc_path.read_bytes(), doc_path.name),
        caption="📄 Tayyor Word fayli",
    )


@dp.message(CommandStart())
async def start_handler(message: types.Message):
    await message.answer(
        f"Assalomu alaykum, {message.from_user.full_name}!\n"
        "Talabalarga va tadqiqotchilarga yordam beruvchi botimizga xush kelibsiz.\n\n"
        "Kerakli bo'limni tanlang:",
        reply_markup=main_menu,
    )


@dp.message(F.text == "📚 Xizmatlar")
async def services_handler(message: types.Message):
    await message.answer("Biz taklif etadigan xizmatlar ro'yxati:", reply_markup=service_menu)


@dp.message(F.text == "ℹ️ Biz haqimizда")
async def about_handler(message: types.Message):
    await message.answer(
        "Talaba servis bot — talabalar va tadqiqotchilar uchun akademik ishlarni tez, sifatli va shaffof tarzda tayyorlashga yordam beradi.\n\n"
        "Biz slaydlar, referatlar, maqolalar, krossvordlar, essay va boshqa akademik matnlar ustida ishlaymiz."
    )


@dp.message(F.text == "✍️ Buyurtma berish")
async def order_start_handler(message: types.Message, state: FSMContext):
    await state.clear()
    await message.answer("Qaysi xizmatni tanlamoqchisiz?", reply_markup=service_menu)
    await state.set_state(OrderForm.waiting_for_service)


@dp.callback_query(F.data.startswith("service_"))
async def process_service_click(callback: types.CallbackQuery, state: FSMContext):
    service = normalize_service_name(callback.data)
    if service not in SERVICES:
        await callback.answer("Noto'g'ri xizmat tanlandi.", show_alert=True)
        return

    await state.update_data(service=service)
    await callback.message.answer(
        f"Siz {SERVICES[service]} xizmatini tanladingiz.\n\n"
        "Endi mavzuingizni yozing. Masalan: 'Sustainable development and innovation in higher education'"
    )
    await state.set_state(OrderForm.waiting_for_topic)
    await callback.answer()


@dp.message(OrderForm.waiting_for_topic)
async def handle_topic(message: types.Message, state: FSMContext):
    topic = message.text.strip()
    if len(topic) < 5:
        await message.answer("Mavzu juda qisqa. Iltimos, kamida 5 ta belgidan iborat mavzu yozing.")
        return

    await state.update_data(topic=topic)
    await message.answer("Nechta slayd kerak bo'ladi? (3 dan 30 gacha son kiriting)")
    await state.set_state(OrderForm.waiting_for_slide_count)


@dp.message(OrderForm.waiting_for_slide_count)
async def handle_slide_count(message: types.Message, state: FSMContext):
    slide_count = parse_slide_count(message.text)
    if slide_count is None:
        await message.answer("Noto'g'ri qiymat kiritdingiz. Iltimos, 3 dan 30 gacha son kiriting.")
        return

    await state.update_data(slide_count=slide_count)
    await message.answer("Tilni tanlang: uz, ru yoki en")
    await state.set_state(OrderForm.waiting_for_language)


@dp.message(OrderForm.waiting_for_language)
async def handle_language(message: types.Message, state: FSMContext):
    language = message.text.strip().lower()
    if language not in {"uz", "ru", "en"}:
        await message.answer("Til noto'g'ri tanlandi. Iltimos: uz, ru yoki en deb yozing.")
        return

    data = await state.get_data()
    service = data.get("service")
    topic = data.get("topic")
    slide_count = data.get("slide_count")

    if not all([service, topic, slide_count]):
        await message.answer("Buyurtma ma'lumotlari to'liq emas. Qayta boshlang.")
        await state.clear()
        return

    await state.update_data(language=language)
    await message.answer("AI matn tayyorlanmoqda. Bir necha soniya kuting...")
    await state.set_state(OrderForm.generating)

    await send_generated_files(message.chat.id, topic, service, slide_count, language)
    await state.clear()


@dp.message(F.text == "📞 Bog'lanish")
async def contact_handler(message: types.Message):
    await message.answer(
        "📞 Biz bilan bog'lanish:\n\n"
        "Savollaringiz, takliflaringiz yoki maxsus buyurtmalaringiz bo'lsa, administratorga murojaat qilishingiz mumkin.\n\n"
        "👤 Admin: @ilhomjon_7oo7\n"
        "📱 Telefon: +998777079749\n"
        "⏰ Ish vaqti: 09:00 - 00:00"
    )


async def main():
    logging.basicConfig(level=logging.INFO)
    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
